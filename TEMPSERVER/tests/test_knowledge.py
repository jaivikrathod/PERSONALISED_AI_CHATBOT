"""Phase 2b — documents as knowledge (WORKFLOW.md B2.2, B4; TASKS.md 2b).

Pinned here:

1. Extraction and chunking: PDF/DOCX/TXT/MD/HTML; chunks fit the embedding
   model's window, never split a sentence, overlap, and cite their section.
2. Ingestion is a job with progress; an unchanged re-upload is a no-op.
3. `questions` remains the FAQ authoring surface and is mirrored into chunks.
4. Hybrid retrieval: a clause number is found even where vectors are weak, and
   nothing weak is returned for an unrelated question.
5. The measurement harness runs, and the shipped gate is measured against it.
"""

import io
import json
from pathlib import Path

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings

from chat.models import ToolExecution
from company.models import Company
from knowledge.chunking import MAX_TOKENS, chunk_blocks
from knowledge.evaluation import evaluate, tune
from knowledge.extraction import Block, ExtractionError, extract
from knowledge.ingest import submit_document
from knowledge.models import IngestionJob, KnowledgeChunk, KnowledgeDocument, KnowledgeSource
from knowledge.retrieval import search
from orchestration.orchestrator import resolve_session, run_turn
from questions.models import Question
from registry.models import Chatbot
from registry.presets import provision_chatbot
from tests.test_orchestration import FakeProvider, add_faq, calls, said
from vector_question.services import EMBEDDING_MODEL_NAME, count_tokens

FIXTURE_PATH = Path(settings.BASE_DIR) / "tests" / "fixtures" / "knowledge_eval.json"
FIXTURE = json.loads(FIXTURE_PATH.read_text())


def assert_same_vector(test, stored, original):
    """pgvector stores float32; compare to that precision."""
    test.assertEqual(len(stored), len(original))
    for a, b in zip(stored, original):
        test.assertAlmostEqual(float(a), float(b), places=5)


def make_company(name="Acme", email="acme@example.com"):
    return Company.objects.create(name=name, email=email, mobile="1", address="a")


def make_pdf(lines):
    """A minimal one-page PDF with real text, built by hand."""
    stream = "BT /F1 12 Tf 72 720 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def load_fixture_tenant(company, chatbot):
    for question, answer in FIXTURE["faqs"]:
        add_faq(company, question, answer)
    source = KnowledgeSource.objects.create(company=company, chatbot=chatbot, name="Policies", kind="document")
    for doc in FIXTURE["documents"]:
        submit_document(source, doc["filename"], doc["text"].encode())
    return source


# --- Extraction ----------------------------------------------------------------------
class ExtractionTests(TestCase):
    def test_markdown_headings_become_sections(self):
        title, blocks = extract("policy.md", b"# Returns\n\nIntro line one\nstill intro.\n\n## Clause 4.2 Damaged\n\nReport it.\n")
        self.assertEqual(title, "Returns")
        self.assertEqual([(b.heading, b.text) for b in blocks], [
            ("Returns", "Intro line one still intro."),
            ("Clause 4.2 Damaged", "Report it."),
        ])

    def test_html_drops_scripts_and_keeps_headings(self):
        html = b"<html><head><title>Shipping</title><script>evil()</script></head><body><h2>Rates</h2><p>Free over 500.</p><ul><li>Fast</li></ul></body></html>"
        title, blocks = extract("shipping.html", html)
        self.assertEqual(title, "Shipping")
        self.assertEqual([(b.heading, b.text) for b in blocks], [("Rates", "Free over 500."), ("Rates", "Fast")])

    def test_docx_paragraphs_and_headings(self):
        import docx

        document = docx.Document()
        document.add_heading("Privacy Policy", level=1)
        document.add_paragraph("We never sell your data.")
        buffer = io.BytesIO()
        document.save(buffer)

        title, blocks = extract("privacy.docx", buffer.getvalue())
        self.assertEqual(title, "Privacy Policy")
        self.assertEqual([(b.heading, b.text) for b in blocks], [("Privacy Policy", "We never sell your data.")])

    def test_pdf_text_carries_its_page(self):
        _, blocks = extract("terms.pdf", make_pdf(["Clause 7 Governing law.", "Courts of Ahmedabad."]))
        self.assertIn("Governing law", " ".join(b.text for b in blocks))
        self.assertEqual({b.page for b in blocks}, {1})

    def test_unsupported_and_empty_files_are_rejected(self):
        with self.assertRaises(ExtractionError):
            extract("sheet.xlsx", b"whatever")
        with self.assertRaises(ExtractionError):
            extract("empty.txt", b"   \n\n  ")


class ChunkingTests(TestCase):
    def test_chunks_fit_the_model_never_split_sentences_and_overlap(self):
        sentences = [f"Sentence number {i} explains one more detail about returns and refunds." for i in range(80)]
        blocks = [Block(" ".join(sentences[i : i + 4]), "Policy") for i in range(0, 80, 4)]

        pieces = chunk_blocks(blocks, count_tokens)

        self.assertGreater(len(pieces), 3)
        for piece in pieces:
            # The heading is prefixed at ingestion, so it shares the budget.
            self.assertLessEqual(count_tokens(f"{piece.section_heading}\n\n{piece.content}"), MAX_TOKENS)
            self.assertTrue(piece.content.endswith("."), piece.content[-40:])
            self.assertEqual(piece.section_heading, "Policy")
        for earlier, later in zip(pieces, pieces[1:]):
            last_sentence = earlier.content.rsplit(". ", 1)[-1]
            self.assertIn(last_sentence, later.content, "consecutive chunks should overlap")

    def test_a_new_heading_starts_a_new_chunk(self):
        pieces = chunk_blocks([Block("Short one.", "A"), Block("Short two.", "B")], count_tokens)
        self.assertEqual([(p.section_heading, p.content) for p in pieces], [("A", "Short one."), ("B", "Short two.")])


# --- Ingestion jobs -----------------------------------------------------------------
@override_settings(INGESTION_MODE="inline")
class IngestionTests(TestCase):
    policy = b"# Return Policy\n\n## Clause 4.1 Window\n\nReturn within 30 days.\n\n## Clause 4.2 Damaged\n\nReport damage within 48 hours.\n"

    def setUp(self):
        self.company = make_company()
        self.chatbot = provision_chatbot(self.company)
        self.source = KnowledgeSource.objects.create(company=self.company, chatbot=self.chatbot, name="Policies", kind="document")

    def test_a_document_becomes_cited_chunks_through_a_job(self):
        document, job = submit_document(self.source, "returns.md", self.policy)

        job.refresh_from_db()
        self.assertEqual((job.status, job.progress), (IngestionJob.Status.DONE, 100))
        self.assertEqual(job.payload.tobytes() if hasattr(job.payload, "tobytes") else bytes(job.payload), b"")
        document.refresh_from_db()
        self.assertEqual((document.status, document.title), ("ready", "Return Policy"))

        chunks = list(KnowledgeChunk.objects.filter(document=document))
        self.assertEqual([c.metadata["section_heading"] for c in chunks], ["Clause 4.1 Window", "Clause 4.2 Damaged"])
        self.assertTrue(all(c.embedding_model == EMBEDDING_MODEL_NAME and c.embed_text == c.content for c in chunks))

    def test_an_unchanged_reupload_is_a_no_op(self):
        submit_document(self.source, "returns.md", self.policy)
        first_ids = set(KnowledgeChunk.objects.values_list("id", flat=True))

        _, job = submit_document(self.source, "returns.md", self.policy)

        self.assertIsNone(job)
        self.assertEqual(IngestionJob.objects.count(), 1)
        self.assertEqual(set(KnowledgeChunk.objects.values_list("id", flat=True)), first_ids)

    def test_a_changed_reupload_replaces_the_chunks(self):
        submit_document(self.source, "returns.md", self.policy)
        submit_document(self.source, "returns.md", self.policy.replace(b"30 days", b"45 days"))

        self.assertEqual(KnowledgeDocument.objects.count(), 1)
        text = " ".join(KnowledgeChunk.objects.values_list("content", flat=True))
        self.assertIn("45 days", text)
        self.assertNotIn("30 days", text)

    def test_a_failed_extraction_is_a_failed_job_not_a_crash(self):
        document, job = submit_document(self.source, "blank.txt", b"   ")
        job.refresh_from_db()
        document.refresh_from_db()
        self.assertEqual(job.status, IngestionJob.Status.FAILED)
        self.assertIn("No text found", job.error)
        self.assertEqual(document.status, "failed")

    @override_settings(INGESTION_MODE="worker")
    def test_worker_mode_leaves_the_job_for_the_worker(self):
        _, job = submit_document(self.source, "returns.md", self.policy)
        self.assertEqual(job.status, IngestionJob.Status.QUEUED)

        call_command("run_ingestion_worker", "--once", stdout=io.StringIO())

        job.refresh_from_db()
        self.assertEqual(job.status, IngestionJob.Status.DONE)
        self.assertEqual(KnowledgeChunk.objects.filter(source=self.source).count(), 2)


@override_settings(INGESTION_MODE="inline")
class KnowledgeApiTests(TestCase):
    def register(self, name, email):
        body = self.client.post(
            "/api/auth/register/",
            {
                "company": {"name": name, "email": email, "mobile": "1", "address": "a"},
                "admin": {"name": "Ada", "email": f"admin-{email}", "password": "secret123", "gender": "Female", "dob": "1990-01-01"},
            },
            content_type="application/json",
        ).json()
        return {"HTTP_AUTHORIZATION": f"Bearer {body['token']}"}

    def test_upload_then_poll_the_job(self):
        auth = self.register("Acme", "a@example.com")
        source = self.client.post("/api/knowledge-sources/", {"name": "Policies", "kind": "document"}, content_type="application/json", **auth)
        self.assertEqual(source.status_code, 201, source.content)
        source_id = source.json()["id"]

        upload = SimpleUploadedFile("returns.md", b"# Returns\n\nReturn within 30 days.\n")
        response = self.client.post(f"/api/knowledge-sources/{source_id}/upload/", {"file": upload}, **auth)
        self.assertEqual(response.status_code, 202, response.content)

        job = self.client.get(f"/api/ingestion-jobs/{response.json()['job']['id']}/", **auth).json()
        self.assertEqual((job["status"], job["progress"]), ("done", 100))
        documents = self.client.get(f"/api/knowledge-sources/{source_id}/documents/", **auth).json()
        self.assertEqual(documents[0]["chunk_count"], 1)

        again = self.client.post(f"/api/knowledge-sources/{source_id}/upload/", {"file": SimpleUploadedFile("returns.md", b"# Returns\n\nReturn within 30 days.\n")}, **auth)
        self.assertEqual((again.status_code, again.json()["unchanged"]), (200, True))

    def test_rejected_inputs(self):
        auth = self.register("Acme", "a@example.com")
        self.assertEqual(self.client.post("/api/knowledge-sources/", {"name": "Site", "kind": "url"}, content_type="application/json", **auth).status_code, 400)
        source_id = self.client.post("/api/knowledge-sources/", {"name": "P", "kind": "document"}, content_type="application/json", **auth).json()["id"]
        bad = self.client.post(f"/api/knowledge-sources/{source_id}/upload/", {"file": SimpleUploadedFile("sheet.xlsx", b"x")}, **auth)
        self.assertEqual(bad.status_code, 400)

    def test_another_company_cannot_see_the_source(self):
        auth = self.register("Acme", "a@example.com")
        rival = self.register("Rival", "r@example.com")
        source_id = self.client.post("/api/knowledge-sources/", {"name": "P", "kind": "document"}, content_type="application/json", **auth).json()["id"]
        self.assertEqual(self.client.get(f"/api/knowledge-sources/{source_id}/", **rival).status_code, 404)


# --- FAQ mirror -----------------------------------------------------------------------
class FaqMirrorTests(TestCase):
    def setUp(self):
        self.company = make_company()
        self.chatbot = provision_chatbot(self.company)

    def chunk(self, question):
        return KnowledgeChunk.objects.filter(metadata__question_id=question.id).first()

    def test_questions_stay_the_authoring_surface_and_are_mirrored(self):
        question = Question.objects.create(company=self.company, question="Do you ship to Germany?", answer="Yes.")
        chunk = self.chunk(question)
        self.assertEqual((chunk.embed_text, chunk.content, chunk.embedding), ("Do you ship to Germany?", "Yes.", None))
        self.assertEqual(chunk.source.kind, "faq")

        question = add_faq(self.company, "How do refunds work?", "Within 7 days.")
        assert_same_vector(self, self.chunk(question).embedding, question.embedding)
        self.assertEqual(self.chunk(question).embedding_model, EMBEDDING_MODEL_NAME)

        question.answer = "Within 10 days."
        question.save()
        self.assertEqual((self.chunk(question).content, self.chunk(question).embedding), ("Within 10 days.", None))

        question.is_archived = True
        question.save()
        self.assertIsNone(self.chunk(question))

    def test_deleting_a_question_removes_its_chunk(self):
        question = Question.objects.create(company=self.company, question="Q?", answer="A.")
        question.delete()
        self.assertFalse(KnowledgeChunk.objects.exists())

    def test_the_backfill_migration_copies_existing_questions(self):
        question = add_faq(self.company, "What are your hours?", "9 to 6.")
        KnowledgeSource.objects.all().delete()

        import importlib

        from django.apps import apps

        importlib.import_module("knowledge.migrations.0002_backfill_faq_chunks").forwards(apps, None)

        chunk = self.chunk(question)
        self.assertEqual((chunk.embed_text, chunk.content), ("What are your hours?", "9 to 6."))
        assert_same_vector(self, chunk.embedding, question.embedding)

    def test_a_chatbot_created_later_gets_the_existing_faqs(self):
        other = make_company("Later", "l@example.com")
        add_faq(other, "Is parking free?", "Yes.")
        bot = provision_chatbot(other)
        self.assertEqual(KnowledgeChunk.objects.filter(chatbot=bot).count(), 1)


# --- Hybrid retrieval and the harness ---------------------------------------------------------
@override_settings(INGESTION_MODE="inline", ORCHESTRATION_TOOL_THREADS=False)
class RetrievalTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.company = make_company()
        cls.chatbot = provision_chatbot(cls.company)
        with override_settings(INGESTION_MODE="inline"):
            load_fixture_tenant(cls.company, cls.chatbot)

    def test_a_paraphrased_faq_is_accepted_semantically(self):
        retrieval = search(self.chatbot, "can I get my money back")
        self.assertEqual(retrieval.accepted_by, "semantic")
        self.assertEqual(retrieval.chunks[0]["title"], "What is your refund policy?")

    def test_a_clause_number_is_found_by_full_text(self):
        retrieval = search(self.chatbot, "what does clause 4.2 say")
        self.assertIsNotNone(retrieval.accepted_by)
        self.assertEqual(retrieval.chunks[0]["section"], "Clause 4.2 Damaged items")
        self.assertIn("48 hours", retrieval.chunks[0]["content"])

    def test_full_text_alone_can_accept_a_specific_token(self):
        retrieval = search(self.chatbot, "clause 4.3", overrides={"accept_threshold": 0.99})
        self.assertEqual(retrieval.accepted_by, "lexical")
        self.assertEqual(retrieval.chunks[0]["section"], "Clause 4.3 Return shipping")

    def test_an_unrelated_question_returns_nothing_rather_than_weak_chunks(self):
        retrieval = search(self.chatbot, "who painted the Mona Lisa")
        self.assertEqual(retrieval.chunks, [])
        self.assertIsNotNone(retrieval.rejected_by)

    def test_top_k_and_the_token_ceiling_are_respected(self):
        self.assertLessEqual(len(search(self.chatbot, "returns and refunds", top_k=2).chunks), 2)
        capped = search(self.chatbot, "returns and refunds", top_k=10, overrides={"max_context_tokens": 1})
        self.assertEqual(len(capped.chunks), 1)

    def test_another_tenants_chunks_are_never_returned(self):
        other = provision_chatbot(make_company("Other", "o@example.com"))
        self.assertEqual(search(other, "can I get my money back").chunks, [])

    def test_the_shipped_gate_is_measured_against_the_harness(self):
        report = evaluate(self.chatbot, FIXTURE["cases"])
        self.assertEqual(report.not_set_precision, 1.0, report.as_dict()["failures"])
        self.assertGreaterEqual(report.recall, 0.8, report.as_dict()["failures"])

    def test_tuning_meets_the_precision_target(self):
        report = tune(self.chatbot, FIXTURE["cases"][::3], target_precision=1.0)
        self.assertIsNotNone(report)
        self.assertEqual(report.not_set_precision, 1.0)

    def test_the_command_writes_tuned_values_to_the_policy(self):
        out = io.StringIO()
        call_command("evaluate_knowledge", "--chatbot", self.chatbot.slug, "--fixture", str(FIXTURE_PATH), "--tune", "--write", stdout=out)
        self.chatbot.refresh_from_db()
        self.assertIn("accept_threshold", self.chatbot.policy)
        self.assertIn("not_set_precision", out.getvalue())

    def test_search_knowledge_returns_cited_document_passages_in_a_turn(self):
        session = resolve_session(self.company.id)
        run_turn(session, "damaged item?", provider=FakeProvider(
            calls(("search_knowledge", {"query": "my item arrived damaged, what should I do"})), said("Report it within 48 hours."),
        ))
        execution = ToolExecution.objects.get(conversation=session)
        self.assertEqual(execution.status, "ok")
        self.assertIn(execution.compiled_query["accepted_by"], ("semantic", "lexical"))
