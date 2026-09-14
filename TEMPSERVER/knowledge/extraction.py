"""Text extraction for PDF / DOCX / TXT / MD / HTML.

Returns *blocks* — paragraphs tagged with their section heading and page — so
the chunker can split on headings and paragraph boundaries and each chunk can
cite where it came from.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import PurePath

SUPPORTED_TYPES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".html": "html",
    ".htm": "html",
}

MARKDOWN_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")


class ExtractionError(Exception):
    pass


@dataclass
class Block:
    text: str
    heading: str = ""
    page: int | None = None


def file_type_for(filename: str) -> str | None:
    return SUPPORTED_TYPES.get(PurePath(filename).suffix.lower())


def extract(filename: str, content: bytes) -> tuple[str, list[Block]]:
    """Returns (title, blocks). Raises ExtractionError."""
    kind = file_type_for(filename)
    if kind is None:
        raise ExtractionError(
            f"Unsupported file type. Upload one of: {', '.join(sorted(SUPPORTED_TYPES))}."
        )
    title, blocks = {
        "pdf": _pdf,
        "docx": _docx,
        "txt": _plain,
        "md": _plain,
        "html": _html,
    }[kind](content)
    blocks = [b for b in blocks if b.text.strip()]
    if not blocks:
        raise ExtractionError(
            "No text found. Scanned PDFs are images and need OCR, which is not supported."
        )
    return (title or PurePath(filename).stem).strip()[:500], blocks


def raw_text(blocks: list[Block]) -> str:
    lines, heading = [], None
    for block in blocks:
        if block.heading and block.heading != heading:
            lines.append(f"## {block.heading}")
            heading = block.heading
        lines.append(block.text)
    return "\n\n".join(lines)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode the file as text.")


def _paragraphs(text: str) -> list[str]:
    # Hard-wrapped lines inside a paragraph are one paragraph.
    return [re.sub(r"\s*\n\s*", " ", p).strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def _plain(content: bytes):
    blocks, heading, title, buffer = [], "", "", []

    def flush():
        for paragraph in _paragraphs("\n".join(buffer)):
            blocks.append(Block(paragraph, heading))
        buffer.clear()

    for line in _decode(content).splitlines():
        match = MARKDOWN_HEADING.match(line)
        if match:
            flush()
            heading = match.group(1)
            title = title or heading
        else:
            buffer.append(line)
    flush()
    return title, blocks


def _pdf(content: bytes):
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        title = (reader.metadata.title if reader.metadata else "") or ""
        blocks = []
        for number, page in enumerate(reader.pages, start=1):
            for paragraph in _paragraphs(page.extract_text() or ""):
                blocks.append(Block(paragraph, "", number))
    except PdfReadError as exc:
        raise ExtractionError(f"Could not read the PDF: {exc}") from None
    return title, blocks


def _docx(content: bytes):
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(content))
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        raise ExtractionError(f"Could not read the DOCX file: {exc}") from None

    title = document.core_properties.title or ""
    blocks, heading = [], ""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.startswith("Heading") or style == "Title":
            heading = text
            title = title or text
        else:
            blocks.append(Block(text, heading))
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(Block(" | ".join(cells), heading))
    return title, blocks


class _HtmlText(HTMLParser):
    HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
    BREAKS = {"p", "div", "li", "br", "tr", "section", "article", "table", "ul", "ol"}
    SKIP = {"script", "style", "noscript", "template", "svg"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.title = ""
        self.heading = ""
        self.buffer: list[str] = []
        self.skipping = 0
        self.in_heading = False
        self.in_title = False

    def _flush(self):
        text = re.sub(r"\s+", " ", "".join(self.buffer)).strip()
        self.buffer = []
        if not text:
            return
        if self.in_heading:
            self.heading = text
        elif not self.in_title:
            self.blocks.append(Block(text, self.heading))

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif tag in self.HEADINGS or tag in self.BREAKS:
            self._flush()
            self.in_heading = tag in self.HEADINGS
        elif tag == "title":
            self._flush()
            self.in_title = True

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skipping = max(0, self.skipping - 1)
        elif tag in self.HEADINGS or tag in self.BREAKS:
            self._flush()
            self.in_heading = False
        elif tag == "title":
            self.title = re.sub(r"\s+", " ", "".join(self.buffer)).strip()
            self.buffer = []
            self.in_title = False

    def handle_data(self, data):
        if not self.skipping:
            self.buffer.append(data)


def _html(content: bytes):
    parser = _HtmlText()
    parser.feed(_decode(content))
    parser._flush()
    return parser.title, parser.blocks
